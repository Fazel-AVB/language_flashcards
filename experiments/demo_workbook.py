"""Build a throwaway vocabulary workbook for the demo notebooks.

WHY THIS EXISTS: the real reference workbook lives in `data/`, which is
gitignored -- your vocabulary is yours. The notebooks in this directory therefore
generate their own 48-word workbook so they run on a fresh clone with no setup.

It deliberately reproduces the quirks the app has to cope with:
  * the column D/E header *labels* differ per topic sheet (Article vs
    Article / Auxiliary vs Article (n/a)), while the column *positions* are fixed
  * a formula-driven dashboard sheet that must never be parsed or rewritten
  * a spread of confidence scores including ungraded (blank) words

Do not "tidy" the header variants into one shared list: they are the whole point
of `models.parse_row` mapping by position rather than by label.
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

# Column A..L, matching models.FIELD_ORDER. Only column D and E labels vary.
_BASE_HEADERS = [
    "German", "English", "Word type", "Article", "Plural", "Pronunciation",
    "Example (German)", "Example (English)", "Synonyms / Related", "Notes",
    "Date added", "Confidence (1-5)",
]
_NOUN_HEADERS = list(_BASE_HEADERS)
_VERB_HEADERS = _BASE_HEADERS[:3] + ["Article / Auxiliary", "Plural / Verb forms"] + _BASE_HEADERS[5:]
_ADJ_HEADERS = _BASE_HEADERS[:3] + ["Article (n/a)", "Comparative / Superlative"] + _BASE_HEADERS[5:]

# (german, english, article/aux, plural/forms, example_de, example_en)
_TOPICS: dict[str, tuple[list[str], str, list[tuple[str, ...]]]] = {
    "School & Office": (_NOUN_HEADERS, "Noun", [
        ("Stift", "pen", "der", "die Stifte", "Ich schreibe mit einem Stift.", "I write with a pen."),
        ("Bericht", "report", "der", "die Berichte", "Ich schreibe einen Bericht.", "I am writing a report."),
        ("Termin", "appointment", "der", "die Termine", "Ich habe morgen einen Termin.", "I have an appointment tomorrow."),
        ("Besprechung", "meeting", "die", "die Besprechungen", "Die Besprechung dauert eine Stunde.", "The meeting lasts an hour."),
        ("Unterlagen", "documents", "die", "nur Plural", "Bringen Sie bitte die Unterlagen mit.", "Please bring the documents."),
        ("Schreibtisch", "desk", "der", "die Schreibtische", "Mein Schreibtisch ist unordentlich.", "My desk is messy."),
        ("Drucker", "printer", "der", "die Drucker", "Der Drucker ist kaputt.", "The printer is broken."),
        ("Frist", "deadline", "die", "die Fristen", "Die Frist laeuft morgen ab.", "The deadline expires tomorrow."),
    ]),
    "Household": (_NOUN_HEADERS, "Noun", [
        ("Waschmaschine", "washing machine", "die", "die Waschmaschinen", "Die Waschmaschine laeuft.", "The washing machine is running."),
        ("Geschirr", "dishes", "das", "kein Plural", "Ich spuele das Geschirr.", "I am washing the dishes."),
        ("Staubsauger", "vacuum cleaner", "der", "die Staubsauger", "Der Staubsauger ist laut.", "The vacuum cleaner is loud."),
        ("Kuehlschrank", "fridge", "der", "die Kuehlschraenke", "Die Milch ist im Kuehlschrank.", "The milk is in the fridge."),
        ("Handtuch", "towel", "das", "die Handtuecher", "Gib mir bitte ein Handtuch.", "Please give me a towel."),
        ("Muelleimer", "bin", "der", "die Muelleimer", "Der Muelleimer ist voll.", "The bin is full."),
        ("Steckdose", "socket", "die", "die Steckdosen", "Wo ist die naechste Steckdose?", "Where is the nearest socket?"),
        ("Vorhang", "curtain", "der", "die Vorhaenge", "Zieh bitte den Vorhang zu.", "Please close the curtain."),
    ]),
    "Travel & Transport": (_NOUN_HEADERS, "Noun", [
        ("Bahnhof", "train station", "der", "die Bahnhoefe", "Wir treffen uns am Bahnhof.", "We are meeting at the station."),
        ("Fahrkarte", "ticket", "die", "die Fahrkarten", "Ich kaufe eine Fahrkarte.", "I am buying a ticket."),
        ("Gepaeck", "luggage", "das", "kein Plural", "Mein Gepaeck ist schwer.", "My luggage is heavy."),
        ("Verspaetung", "delay", "die", "die Verspaetungen", "Der Zug hat Verspaetung.", "The train is delayed."),
        ("Haltestelle", "stop", "die", "die Haltestellen", "Die Haltestelle ist um die Ecke.", "The stop is around the corner."),
        ("Flugzeug", "aeroplane", "das", "die Flugzeuge", "Das Flugzeug startet gleich.", "The plane is about to take off."),
        ("Unterkunft", "accommodation", "die", "die Unterkuenfte", "Die Unterkunft war guenstig.", "The accommodation was cheap."),
        ("Umstieg", "change (of train)", "der", "die Umstiege", "Der Umstieg ist knapp.", "The connection is tight."),
    ]),
    "Verbs": (_VERB_HEADERS, "Verb", [
        ("werden", "to become", "sein", "wird - wurde - ist geworden", "Sie will Aerztin werden.", "She wants to become a doctor."),
        ("gelingen", "to succeed", "sein", "gelingt - gelang - ist gelungen", "Der Versuch ist gelungen.", "The attempt succeeded."),
        ("verlassen", "to leave", "haben", "verlaesst - verliess - hat verlassen", "Ich verlasse das Haus.", "I am leaving the house."),
        ("erledigen", "to get done", "haben", "erledigt - erledigte - hat erledigt", "Ich erledige das heute.", "I will get that done today."),
        ("vermeiden", "to avoid", "haben", "vermeidet - vermied - hat vermieden", "Wir vermeiden Fehler.", "We avoid mistakes."),
        ("beitragen", "to contribute", "haben", "traegt bei - trug bei - hat beigetragen", "Jeder traegt etwas bei.", "Everyone contributes something."),
        ("ueberlegen", "to consider", "haben", "ueberlegt - ueberlegte - hat ueberlegt", "Ich ueberlege es mir.", "I will think it over."),
        ("sich beeilen", "to hurry", "haben", "beeilt sich - beeilte sich - hat sich beeilt", "Beeil dich bitte!", "Please hurry up!"),
    ]),
    "Adjectives & Adverbs": (_ADJ_HEADERS, "Adjective / Adverb", [
        ("riesig", "huge", "-", "riesiger - am riesigsten", "ein riesiges Haus", "a huge house"),
        ("zuverlaessig", "reliable", "-", "zuverlaessiger - am zuverlaessigsten", "ein zuverlaessiger Kollege", "a reliable colleague"),
        ("anspruchsvoll", "demanding", "-", "anspruchsvoller - am anspruchsvollsten", "eine anspruchsvolle Aufgabe", "a demanding task"),
        ("vorlaeufig", "provisional", "-", "-", "ein vorlaeufiges Ergebnis", "a provisional result"),
        ("ausfuehrlich", "detailed", "-", "ausfuehrlicher - am ausfuehrlichsten", "ein ausfuehrlicher Bericht", "a detailed report"),
        ("knapp", "scarce, tight", "-", "knapper - am knappsten", "Die Zeit ist knapp.", "Time is tight."),
        ("allmaehlich", "gradual", "-", "-", "Es wird allmaehlich kaelter.", "It is gradually getting colder."),
        ("geeignet", "suitable", "-", "geeigneter - am geeignetsten", "ein geeigneter Ort", "a suitable place"),
    ]),
    "Feelings & Abstract": (_NOUN_HEADERS, "Noun", [
        ("Vertrauen", "trust", "das", "kein Plural", "Vertrauen muss man verdienen.", "Trust has to be earned."),
        ("Zweifel", "doubt", "der", "die Zweifel", "Ich habe da Zweifel.", "I have doubts about that."),
        ("Erwartung", "expectation", "die", "die Erwartungen", "Die Erwartungen sind hoch.", "Expectations are high."),
        ("Enttaeuschung", "disappointment", "die", "die Enttaeuschungen", "Das war eine Enttaeuschung.", "That was a disappointment."),
        ("Gelassenheit", "composure", "die", "kein Plural", "Er reagiert mit Gelassenheit.", "He reacts with composure."),
        ("Neugier", "curiosity", "die", "kein Plural", "Aus reiner Neugier.", "Out of pure curiosity."),
        ("Ehrgeiz", "ambition", "der", "kein Plural", "Ihm fehlt der Ehrgeiz.", "He lacks ambition."),
        ("Sehnsucht", "longing", "die", "die Sehnsuechte", "Sehnsucht nach Hause.", "A longing for home."),
    ]),
}

# Score mix: a realistic library is mostly half-learned, with a few never seen.
# `None` is an ungraded word. Drawn with a fixed seed so every notebook run, and
# every reader, sees the same numbers.
_SCORE_POOL = [None, None, 1, 1, 2, 2, 2, 3, 3, 3, 3, 4, 4, 4, 5, 5]


def make_demo_workbook(path: str | Path, seed: int = 7) -> Path:
    """Write a demo .xlsx to `path` and return the path.

    Overwrites `path` if it exists -- the notebooks rebuild it on every run so a
    half-finished write-back from a previous run never leaks into the next.
    """
    rng = random.Random(seed)
    start = date(2026, 6, 1)

    wb = Workbook()

    # The dashboard sheet goes first, as in the real workbook. It is formula-driven
    # and `models.NON_TOPIC_SHEETS` makes the app skip it entirely; notebook 03
    # checks it survives a write-back byte for byte.
    dash = wb.active
    dash.title = "Overview"
    dash["A1"] = "Topic"
    dash["B1"] = "Words"
    dash["C1"] = "Average confidence"
    for i, topic in enumerate(_TOPICS, start=2):
        dash[f"A{i}"] = topic
        dash[f"B{i}"] = f"=COUNTA('{topic}'!A2:A1000)"
        # The real workbook uses LET/LAMBDA dynamic arrays here; openpyxl cannot
        # emit those, so a plain AVERAGE stands in. The preservation argument is
        # the same either way: openpyxl must never re-save this file.
        dash[f"C{i}"] = f"=IFERROR(AVERAGE('{topic}'!L2:L1000),\"\")"

    for topic, (headers, word_type, rows) in _TOPICS.items():
        ws = wb.create_sheet(topic)
        ws.append(headers)
        for n, (german, english, article, forms, ex_de, ex_en) in enumerate(rows):
            ws.append([
                german, english, word_type, article, forms,
                "",                                   # pronunciation: often blank in practice
                ex_de, ex_en,
                "", "",                               # synonyms / notes: often blank
                start + timedelta(days=rng.randrange(0, 90)),
                rng.choice(_SCORE_POOL),
            ])

    # A second dashboard-style sheet, to show name-based skipping as well.
    pool = wb.create_sheet("Card Pool")
    pool.append(["Topic", "German", "English"])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":  # quick manual check: python demo_workbook.py
    p = make_demo_workbook("_demo/demo_vocab.xlsx")
    print(f"wrote {p}")
