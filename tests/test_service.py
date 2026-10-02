"""service.Service: the full path from choosing a workbook to saving a session."""

import csv
import os
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest

from flashcards import config, paths
from flashcards.service import Service, UserError

T0 = datetime(2026, 10, 2, 9, 0, tzinfo=timezone.utc)


class Clock:
    def __init__(self, t):
        self.t = t

    def __call__(self):
        return self.t


@pytest.fixture
def clock():
    return Clock(T0)


@pytest.fixture
def svc(workbook_path, clock):
    s = Service(now=clock, local_tz=timezone.utc, fuzz=False)
    s.set_reference(str(workbook_path))
    return s


DUE = {"mode": "due", "size": 20}


def run_session(svc, grade=4, params=DUE):
    view = svc.start(params)
    while view["stage"] == "review":
        view = svc.grade(grade)
    return view


def test_overview_before_and_after_reference(workbook_path, clock):
    s = Service(now=clock)
    assert s.overview()["reference"] is None
    o = s.set_reference(str(workbook_path))
    assert o["words"] == 4 and o["new"] == 4 and o["due"] == 0
    assert o["languages"] == {"term": "German", "native": "English"}
    assert [t["topic"] for t in o["topics"]] == ["School & Office", "Verbs", "Adjectives & Adverbs"]


def test_bad_reference_is_user_error(tmp_path, clock):
    with pytest.raises(UserError, match="Not a file"):
        Service(now=clock).set_reference(str(tmp_path / "nope.xlsx"))


def test_full_session_saves_confidence_schedule_and_log(svc, workbook_path, clock):
    summary = run_session(svc, grade=4)
    assert summary["stage"] == "summary" and len(summary["rows"]) == 4

    saved = svc.write()
    assert saved["stage"] == "saved" and os.path.isfile(saved["backup"])
    assert svc.view()["stage"] == "saved"

    ws = openpyxl.load_workbook(workbook_path)["School & Office"]
    assert ws["M1"].value == "Due (FSRS)"           # headers added on first save
    assert ws["N2"].value == "Learning"             # first Good on a new word
    assert ws["L2"].value == 2.5                     # Stift: (1 + 4) / 2

    o = svc.overview()
    assert o["new"] == 0 and o["later"] == 4 and o["due"] == 0

    with open(paths.reviews_path(), newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 4 and {r["fsrs_rating"] for r in rows} == {"Good"}

    # Ten minutes on, the learning step is due again.
    clock.t = T0 + timedelta(minutes=11)
    assert svc.overview()["due"] == 4


def test_card_view_has_previews(svc):
    view = svc.start({"mode": "mix", "size": 1, "direction": "term_to_native"})
    assert view["stage"] == "review"
    assert [g["grade"] for g in view["grades"]] == [1, 2, 3, 4, 5]
    assert view["grades"][2]["next"] == "next review in 10 min"   # Good on a new word
    assert view["card"]["schedule"] == "new word"
    assert view["card"]["front"]["text"] == view["card"]["back"]["term"]
    # Language codes for the browser's speech voices (app.js voiceFor).
    assert view["card"]["front"]["lang"] == "de"
    assert view["card"]["back"]["native_lang"] == "en"


def test_prefs_are_remembered(svc):
    svc.start({"mode": "topic", "topics": ["Verbs"], "size": 5, "repeat": True})
    prefs = config.load_config()["prefs"]
    assert prefs["mode"] == "topic" and prefs["topics"] == ["Verbs"] and prefs["session_size"] == 5


def test_workbook_changed_during_session(svc, workbook_path):
    run_session(svc)
    os.utime(workbook_path, (1, 1))
    with pytest.raises(UserError, match="changed on disk"):
        svc.write()
    assert svc.view()["stage"] == "summary"  # nothing lost; you can discard or retry


def test_end_early_and_discard(svc, workbook_path):
    before = workbook_path.read_bytes()
    svc.start(DUE)
    svc.grade(3)
    assert svc.end()["stage"] == "summary"
    assert svc.discard()["stage"] == "setup"
    assert workbook_path.read_bytes() == before


@pytest.mark.parametrize("params,msg", [
    ({"mode": "bogus"}, "Unknown mode"),
    ({"mode": "topic"}, "at least one topic"),
    ({"mode": "mix", "size": 0}, "between 1"),
    ({"mode": "mix", "threshold": "x"}, "whole numbers"),
    ({"mode": "difficult", "threshold": 1, "topics": ["Verbs"]}, "No words"),
])
def test_bad_deck_settings(svc, params, msg):
    with pytest.raises(UserError, match=msg):
        svc.start(params)


def test_one_session_at_a_time(svc, workbook_path):
    svc.start(DUE)
    with pytest.raises(UserError, match="already"):
        svc.start(DUE)
    with pytest.raises(UserError, match="Finish or discard"):
        svc.set_reference(str(workbook_path))


def test_grade_validation(svc):
    svc.start(DUE)
    for bad in (0, 6, "x", None):
        with pytest.raises(UserError):
            svc.grade(bad)


def test_scope_counts(svc):
    s = svc.scope({"mode": "due", "size": 2})
    assert s == {"in_scope": 4, "drawn": 2, "due": 0, "new": 2}


def test_article_deck_has_only_der_die_das_nouns(svc):
    # Fixture: Stift and Bericht are "der" nouns; werden has "sein", riesig "—".
    assert svc.overview()["article_words"] == 2
    assert svc.scope({"mode": "mix", "exercise": "article"})["in_scope"] == 2
    view = svc.start({"mode": "mix", "exercise": "article", "direction": "native_to_term"})
    assert view["exercise"]["kind"] == "article"
    assert view["card"]["front"]["text"] == view["card"]["back"]["term"]   # the noun, whatever the direction


def test_recall_shows_the_meaning(svc):
    view = svc.start({"mode": "mix", "size": 1, "exercise": "recall", "direction": "term_to_native"})
    assert view["exercise"]["kind"] == "recall"
    assert view["card"]["front"]["text"] == view["card"]["back"]["meaning"]
    assert svc.check(view["card"]["back"]["term"])["result"]["verdict"] == "correct"


def test_listening_needs_the_browser_voice(svc):
    with pytest.raises(UserError, match="needs a German voice"):
        svc.start({"mode": "mix", "exercise": "listen", "voices": ["en"]})
    view = svc.start({"mode": "mix", "size": 1, "exercise": "listen", "voices": ["de", "en"]})
    ex = view["exercise"]
    assert ex["kind"] == "listen" and ex["say"] == view["card"]["back"]["term"] and ex["lang"] == "de"
    assert view["card"]["front"]["text"] == ""          # the word is heard, not shown
    assert svc.check(ex["say"])["result"]["verdict"] == "correct"


def test_mixed_without_voice_never_listens(svc):
    view = svc.start({"mode": "mix", "exercise": "mixed", "voices": []})
    kinds = set()
    while view["stage"] == "review":
        kinds.add(view["exercise"]["kind"])
        view = svc.skip()
    assert "listen" not in kinds
