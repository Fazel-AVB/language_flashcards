"""Read the reference workbook into a list of Cards (read-only).

Only *topic* sheets are read (see `models.is_topic_sheet`). The dashboard sheets
are formula-driven and are never parsed or written by this app.

The workbook declares its own language pair in the A1/B1 headers of each topic
sheet, so this module works unchanged for German/English, English/Persian, or any
other pair. `WorkbookData.languages` is the pair the UI labels itself with.

Downstream: `WorkbookData.mtime` is captured here and re-checked by
`writeback.apply_updates` -- if the file changed on disk between load
and save, the write is refused so a hand edit in Excel is never clobbered.
"""

from __future__ import annotations

import dataclasses
import os
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import load_workbook

from .models import (
    CONFIDENCE_COL_INDEX,
    Card,
    LanguagePair,
    is_topic_sheet,
    language_pair_from_header,
    parse_row,
)
from .srs import SRS_COLUMNS, SRS_FIRST_COL_INDEX, SRS_HEADERS, SRS_LAST_COL_INDEX, schedule_from_cells


class WorkbookError(RuntimeError):
    """The file could not be read as a valid reference workbook."""


@dataclass
class TopicStat:
    topic: str
    count: int
    graded: int
    avg_confidence: float | None
    difficult: int  # cards at or below the difficult threshold


@dataclass
class WorkbookData:
    path: str
    mtime: float
    cards: list[Card]
    languages: LanguagePair
    topics: list[str] = field(default_factory=list)
    # Sheets whose A1/B1 disagree with the first topic sheet's pair. Surfaced as a
    # warning rather than an error: each Card carries its own sheet's pair, so a
    # mixed workbook still reviews correctly, it just cannot be labelled with one
    # pair in the sidebar.
    language_conflicts: list[str] = field(default_factory=list)
    # "Sheet!M7: reason" for every row whose schedule columns could not be read.
    # Such a word is treated as new; listed so a broken hand edit is visible.
    schedule_problems: list[str] = field(default_factory=list)
    # Sheets whose M1..R1 do not yet carry the schedule headers. Write-back adds
    # them, so the first save that schedules a word also labels the columns.
    sheets_missing_srs_headers: list[str] = field(default_factory=list)

    def stats(self, difficult_threshold: int) -> list[TopicStat]:
        out: list[TopicStat] = []
        for topic in self.topics:
            group = [c for c in self.cards if c.topic == topic]
            graded = [c.confidence for c in group if c.confidence is not None]
            out.append(
                TopicStat(
                    topic=topic,
                    count=len(group),
                    graded=len(graded),
                    avg_confidence=(sum(graded) / len(graded)) if graded else None,
                    difficult=sum(
                        1 for c in group
                        if c.confidence is not None and c.confidence <= difficult_threshold
                    ),
                )
            )
        return out


def load_cards(path: str) -> WorkbookData:
    """Parse every topic sheet into Cards.

    Raises WorkbookError with a user-facing message on anything unreadable so the
    UI can show it verbatim.
    """
    if not os.path.isfile(path):
        raise WorkbookError(f"File not found: {path}")

    # Capture mtime BEFORE reading so a concurrent write is caught by the guard,
    # not masked by our own read time.
    mtime = os.path.getmtime(path)

    try:
        # data_only: topic sheets hold no formulas, but this also returns cached
        # values for any that appear later instead of the formula string.
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        raise WorkbookError(f"Could not open workbook: {exc}") from exc

    cards: list[Card] = []
    topics: list[str] = []
    pairs: dict[str, LanguagePair] = {}
    problems: list[str] = []
    missing_headers: list[str] = []
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            try:
                header = list(next(rows))
            except StopIteration:
                continue
            if not is_topic_sheet(name, header):
                continue

            topics.append(name)
            srs_header = (list(header) + [None] * SRS_LAST_COL_INDEX)[SRS_FIRST_COL_INDEX - 1:SRS_LAST_COL_INDEX]
            if [str(h or "").strip() for h in srs_header] != [SRS_HEADERS[c] for c in SRS_COLUMNS]:
                missing_headers.append(name)
            header = header[:CONFIDENCE_COL_INDEX]
            languages = language_pair_from_header(header)
            pairs[name] = languages
            # openpyxl is 1-based; row 1 is the header, data starts at row 2.
            for offset, values in enumerate(rows, start=2):
                values = list(values)
                card = parse_row(
                    name, header, values[:CONFIDENCE_COL_INDEX], offset, languages
                )
                if card is None:
                    continue
                try:
                    schedule = schedule_from_cells(values[SRS_FIRST_COL_INDEX - 1:SRS_LAST_COL_INDEX])
                except (ValueError, TypeError) as exc:
                    problems.append(f"{name}!M{offset}: {exc}")
                    schedule = None
                cards.append(dataclasses.replace(card, schedule=schedule))
    finally:
        wb.close()

    if not topics:
        raise WorkbookError(
            "No topic sheets found. A topic sheet needs the target language in A1, "
            "the native language in B1, and 'Confidence (1-5)' in L1, with the "
            "header on row 1."
        )

    # The first topic sheet defines the workbook's pair; the rest are checked
    # against it so a half-converted workbook is visible instead of silent.
    primary = pairs[topics[0]]
    conflicts = [name for name, pair in pairs.items() if pair != primary]

    return WorkbookData(
        path=path,
        mtime=mtime,
        cards=cards,
        languages=primary,
        topics=topics,
        language_conflicts=conflicts,
        schedule_problems=problems,
        sheets_missing_srs_headers=missing_headers,
    )


def describe_reference(path: str) -> str:
    """One-line summary for the sidebar."""
    try:
        data = load_cards(path)
    except WorkbookError as exc:
        return f"unreadable ({exc})"
    when = datetime.fromtimestamp(data.mtime).strftime("%Y-%m-%d %H:%M")
    pair = f"{data.languages.term}/{data.languages.native}"
    return f"{pair} - {len(data.cards)} words / {len(data.topics)} topics / modified {when}"
